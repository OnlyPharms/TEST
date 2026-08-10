# -*- coding: utf-8 -*-
"""
Buduje plik eksportowy Only Pharms z egzaminu:
Farmakologia, Stomatologia rok III, WUM 2014/2015, wersja 1 (100 pytan).

Uklad kolumn zgodny z JAK_KATEGORYZOWAC.txt i z plikami w questions/*.csv:
Section, Category, Poziom, Type, Question, True 1..5, False 1..5
"""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# (nr, Section, Category, Question, [True...], [False...])
Q = [
(1, "LEKI PRZECIWBAKTERYJNE", "Mechanizm działania",
 "Hamowanie syntezy ściany komórkowej jest mechanizmem działania:",
 ["wankomycyna", "aztreonam", "bacytracyna"],
 ["azytromycyna"]),

(2, "LEKI PRZECIWGRZYBICZE", "Sposób podania",
 "Wskaż właściwe stwierdzenia odnoszące się do leków przeciwgrzybiczych:",
 ["mikonazol może być podawany doustnie",
  "klotrimazol jest stosowany w leczeniu zakażeń wywołanych przez dermatofity",
  "amfoterycyna B może być podawana dożylnie"],
 ["ketokonazol dobrze przechodzi przez barierę krew-mózg"]),

(3, "ZNIECZULENIA MIEJSCOWE", "Działania niepożądane",
 "Do działań niepożądanych lidokainy należy:",
 ["depresja ośrodka oddechowego", "bradykardia",
  "obniżenie ciśnienia tętniczego", "drgawki"],
 ["methemoglobinemia", "hipertermia złośliwa", "agranulocytoza"]),

(4, "ZNIECZULENIA MIEJSCOWE", "Farmakokinetyka",
 "Wskaż właściwe twierdzenia odnoszące się do lidokainy:",
 ["posiada silny efekt pierwszego przejścia", "jest metabolizowana w wątrobie"],
 ["jest typem estrowego leku miejscowo znieczulającego",
  "przedawkowanie wywołuje methemoglobinemię"]),

(5, "LEKI PRZECIWBAKTERYJNE", "Działania niepożądane",
 "Ototoksyczność jest działaniem niepożądanym:",
 ["streptomycyna", "wankomycyna", "furosemid"],
 ["etambutol"]),

(6, "FARMAKOKINETYKA", "Farmakokinetyka",
 "Wskaż właściwe stwierdzenia odnoszące się do kinetyki I rzędu:",
 ["w czasie biologicznego półtrwania stężenie w surowicy zmniejsza się o 50%",
  "w jednostce czasu organizm eliminuje część całkowitego stężenia leku",
  "szybkość reakcji jest wprost proporcjonalna do pierwszej potęgi stężenia"],
 ["eliminacja leku jest stała w jednostce czasu"]),

(7, "LEKI PRZECIWBAKTERYJNE", "Wskazania",
 "W zakażeniu wywołanym przez Pseudomonas aeruginosa skuteczność wykazuje:",
 ["amikacyna", "piperacylina", "lewofloksacyna"],
 ["doksycyklina"]),

(8, "NARKOTYCZNE LEKI PRZECIWBÓLOWE", "Sposób podania",
 "Morfina może być stosowana:",
 ["podskórnie", "podpajęczynówkowo", "dożylnie", "doustnie"],
 ["dospojówkowo", "wyłącznie parenteralnie", "wyłącznie w postaci plastra przezskórnego"]),

(9, "NADCIŚNIENIE TĘTNICZE", "Działania niepożądane",
 "Stężenie lipidów w surowicy ulega zwiększeniu po stosowaniu:",
 ["chlortalidon", "metoprolol"],
 ["doksazosyna", "peryndopryl"]),

(10, "CHOROBA NIEDOKRWIENNA SERCA", "Działania niepożądane",
 "Działaniem niepożądanym nitrogliceryny jest:",
 ["ból głowy", "zaczerwienienie skóry", "działanie hipotensyjne", "częstoskurcz"],
 ["bradykardia", "nadciśnienie tętnicze", "suchość błony śluzowej jamy ustnej"]),

(11, "NADCIŚNIENIE TĘTNICZE", "Działania niepożądane",
 "Niepożądanym działaniem enalaprylu jest:",
 ["działanie hipotensyjne", "zaburzenie smaku", "kaszel", "obrzęk naczynioruchowy"],
 ["hipokaliemia", "hirsutyzm", "przerost dziąseł"]),

(12, "LEKI MOCZOPĘDNE", "Wskazania",
 "Wskazaniem do stosowania spironolaktonu jest:",
 ["niewydolność serca", "wodobrzusze w przebiegu marskości wątroby",
  "pierwotny hiperaldosteronizm"],
 ["choroba wieńcowa"]),

(13, "NARKOTYCZNE LEKI PRZECIWBÓLOWE", "Mechanizm działania",
 "Który z poniższych leków wykazuje działanie agonistyczne w stosunku do receptorów opioidowych?",
 ["metadon", "buprenorfina", "oksykodon", "pentazocyna"],
 ["nalokson", "naltrekson", "flumazenil"]),

(14, "UKŁAD WSPÓŁCZULNY", "Mechanizm działania",
 "Wskaż właściwe stwierdzenia odnoszące się do fentolaminy:",
 ["może zaostrzyć objawy choroby wrzodowej",
  "jest stosowana w leczeniu pheochromocytoma",
  "słabo się wchłania po podaniu doustnym"],
 ["nieodwracalnie blokuje receptory alfa-adrenergiczne"]),

(15, "UKŁAD ODDECHOWY", "Działania niepożądane",
 "Wskaż właściwe stwierdzenia odnoszące się do salbutamolu:",
 ["nasila objawy choroby wieńcowej",
  "wywołuje skurcze włókienkowe mięśni szkieletowych"],
 ["powoduje senność", "obniża stężenie glukozy w surowicy"]),

(16, "UKŁAD ODDECHOWY", "Mechanizm działania",
 "Antagonistą receptorów leukotrienowych jest:",
 ["zafirlukast"],
 ["zileuton", "flumazenil", "dinoprost"]),

(17, "LEKI MOCZOPĘDNE", "Farmakokinetyka",
 "Wskaż właściwe stwierdzenia odnoszące się do torasemidu:",
 ["działa tak samo silnie jak furosemid",
  "w porównaniu z furosemidem ma dłuższy czas biologicznego półtrwania",
  "może być stosowany przewlekle"],
 ["powoduje hiperkalcemię"]),

(18, "LEKI PRZECIWBAKTERYJNE", "Wskazania",
 "Klarytromycyna jest stosowana w zakażeniu:",
 ["Legionella pneumophila", "Chlamydia trachomatis", "Bordetella pertussis"],
 ["Proteus mirabilis"]),

(19, "LEKI PRZECIWGRZYBICZE", "Wskazania",
 "W zakażeniu Epidermophyton możesz zastosować:",
 ["amorolfina", "ketokonazol"],
 ["amfoterycyna B", "flucytozyna"]),

(20, "ANTYKONCEPCJA", "Przeciwwskazania",
 "Przeciwwskazaniem do stosowania doustnych środków antykoncepcyjnych jest:",
 ["choroba naczyniowa mózgu", "nowotwór sutka", "ciąża",
  "choroba zakrzepowo-zatorowa w wywiadzie"],
 ["niedokrwistość z niedoboru żelaza", "trądzik pospolity",
  "zespół policystycznych jajników"]),

(21, "GLIKOKORTYKOSTEROIDY", "Działania niepożądane",
 "Glikokortykosteroidy podane miejscowo do drzewa oskrzelowego mogą spowodować:",
 ["chrypka", "kaszel"],
 ["grzybicze zakażenie zatok przynosowych", "porażenie nerwu krtaniowego"]),

(22, "FARMAKOKINETYKA", "Farmakokinetyka",
 "Efekt pierwszego przejścia zależy od:",
 ["szybkości pasażu przez przewód pokarmowy",
  "szybkości eliminacji leku z organizmu",
  "wielkości metabolizmu leku w wątrobie"],
 ["szybkości wchłaniania"]),

(23, "ZNIECZULENIA OGÓLNE", "Klasyfikacja",
 "Neuroleptanalgezja jest wynikiem łącznego stosowania:",
 ["droperidol i fentanyl"],
 ["fentanyl i propofol", "propofol i podtlenek azotu",
  "droperidol i podtlenek azotu"]),

(24, "LEKI PRZECIWDEPRESYJNE", "Działania niepożądane",
 "Do działań niepożądanych amitryptyliny należy:",
 ["suchość błon śluzowych jamy ustnej", "zaburzenia rytmu serca"],
 ["przerost dziąseł", "zespół parkinsonopodobny"]),

(25, "UKŁAD PRZYWSPÓŁCZULNY", "Wskazania",
 "Atropina jest stosowana w:",
 ["premedykacji", "zatruciu pestycydami", "bradykardii zatokowej"],
 ["leczeniu jaskry"]),

(26, "NIEWYDOLNOŚĆ SERCA", "Wskazania",
 "W leczeniu niewydolności serca możesz zastosować:",
 ["walsartan", "lizynopryl", "digoksyna"],
 ["klopamid"]),

(27, "NADCIŚNIENIE TĘTNICZE", "Wskazania",
 "W przełomie nadciśnieniowym zastosujesz:",
 ["urapidil", "labetalol", "nitroprusydek sodu"],
 ["triamteren"]),

(28, "NADCIŚNIENIE TĘTNICZE", "Przeciwwskazania",
 "Przeciwwskazaniem do podawania bisoprololu jest:",
 ["blok serca", "dychawica oskrzelowa", "bradykardia", "chromanie przestankowe"],
 ["nadciśnienie tętnicze", "częstoskurcz nadkomorowy", "choroba niedokrwienna serca"]),

(29, "UKŁAD WSPÓŁCZULNY", "Działania niepożądane",
 "Wskaż właściwe stwierdzenia odnoszące się do efedryny:",
 ["zaburza akomodację", "wywołuje tachykardię"],
 ["powoduje senność", "obniża ciśnienie tętnicze"]),

(30, "UKŁAD ODDECHOWY", "Mechanizm działania",
 "Wskaż właściwe stwierdzenia odnoszące się do teofiliny:",
 ["jest antagonistą receptorów adenozynowych", "hamuje fosfodiesterazę"],
 ["zmniejsza syntezę prostaglandyn", "hamuje topoizomerazę IV"]),

(31, "LEKI MOCZOPĘDNE", "Działania niepożądane",
 "Hiponatremia może wystąpić w przebiegu stosowania:",
 ["torasemid", "spironolakton", "hydrochlorotiazyd", "karbamazepina"],
 ["demeklocyklina", "tolwaptan", "mannitol"]),

(32, "LEKI PRZECIWBAKTERYJNE", "Działania niepożądane",
 "W przebiegu stosowania ceftazydymu może wystąpić:",
 ["trombocytopenia"],
 ["reakcja disulfiramopodobna", "zespół Hoigne'a", "zespół czerwonego karku"]),

(33, "LEKI PRZECIWBAKTERYJNE", "Działania niepożądane",
 "Wskaż właściwe stwierdzenia odnoszące się do wankomycyny:",
 ["nie wchłania się z przewodu pokarmowego", "działa nefrotoksycznie",
  "zwiększa uwalnianie histaminy",
  "jest stosowana w zakażeniu Clostridium difficile"],
 ["dobrze wchłania się po podaniu doustnym",
  "jest lekiem z wyboru w zakażeniach pałeczkami Gram-ujemnymi",
  "hamuje syntezę białka na rybosomie bakteryjnym"]),

(34, "LEKI PRZECIWBAKTERYJNE", "Wskazania",
 "Wskaż prawidłowe połączenie lek – wskazanie do stosowania:",
 ["kotrimoksazol – bakteryjne zapalenie prostaty",
  "fosfomycyna – niepowikłane zakażenie układu moczowego"],
 ["imipenem – zakażenie gronkowcami metycylinoopornymi",
  "aztreonam – zakażenie bakteriami beztlenowymi"]),

(35, "LEKI PRZECIWBAKTERYJNE", "Wskazania",
 "Penicylinę G zastosujesz w leczeniu:",
 ["angina", "kiła", "tężec", "promienica"],
 ["zakażenie Mycoplasma pneumoniae",
  "zakażenie gronkowcem metycylinoopornym",
  "rzekomobłoniaste zapalenie jelit"]),

(36, "LEKI MOCZOPĘDNE", "Działania niepożądane",
 "Hipokaliemia może wystąpić w przebiegu stosowania:",
 ["chlortalidon", "salmeterol"],
 ["enalapryl", "eplerenon"]),

(37, "GLIKOKORTYKOSTEROIDY", "Wskazania",
 "Glikokortykosteroidy zastosujesz w:",
 ["przełom nadnerczowy", "obrzęk nagłośni", "choroba Addisona",
  "choroba Leśniowskiego-Crohna"],
 ["zespół Cushinga", "osteoporoza pomenopauzalna",
  "jaskra z otwartym kątem przesączania"]),

(38, "LEKI PRZECIWBAKTERYJNE", "Mechanizm działania",
 "Hamowanie syntezy białka komórki bakteryjnej wskutek działania na rybosom wykazuje:",
 ["linezolid", "amikacyna"],
 ["norfloksacyna", "kolistyna"]),

(39, "LEKI PRZECIWBAKTERYJNE", "Działania niepożądane",
 "Metronidazol powoduje:",
 ["metaliczny smak w ustach", "reakcja disulfiramowa"],
 ["uszkodzenie wątroby", "objawy płatowego zapalenia płuc"]),

(40, "TARCZYCA", "Działania niepożądane",
 "Wskaż właściwe stwierdzenia odnoszące się do jodku potasu:",
 ["powoduje obrzęk ślinianek", "może nasilić trądzik",
  "zaburza czynność tarczycy", "jest przeciwwskazany w gruźlicy"],
 ["powoduje agranulocytozę", "jest lekiem pierwszego rzutu w niedoczynności tarczycy",
  "nie przenika do mleka matki"]),

(41, "UKŁAD ODDECHOWY", "Klasyfikacja",
 "Do leków mukolitycznych należy:",
 ["N-acetylocysteina", "ambroksol"],
 ["gwajafenezyna", "okseladyna"]),

(42, "CUKRZYCA", "Mechanizm działania",
 "Wydzielanie insuliny zwiększa:",
 ["glibenklamid", "repaglinid"],
 ["metformina", "pioglitazon"]),

(43, "CUKRZYCA", "Działania niepożądane",
 "Do działań niepożądanych leków inkretynowych – inhibitorów DPP-4 należy:",
 ["zastoinowa niewydolność krążenia", "wzdęcie brzucha",
  "hipoglikemia w skojarzeniu z pochodnymi sulfonylomocznika"],
 ["ostre zapalenie pęcherzyka żółciowego"]),

(44, "TARCZYCA", "Mechanizm działania",
 "Wskaż właściwe stwierdzenia odnoszące się do propylotiouracylu:",
 ["hamuje konwersję T4 w T3", "hamuje syntezę hormonów tarczycy",
  "powoduje agranulocytozę"],
 ["zwiększa ryzyko choroby wrzodowej"]),

(45, "LEKI PRZECIWBAKTERYJNE", "Mechanizm działania",
 "Wskaż prawidłowe połączenie lek przeciwprątkowy – mechanizm działania:",
 ["izoniazyd – hamowanie syntezy kwasów mykolowych wchodzących w skład ściany komórkowej prątka",
  "ryfampicyna – blokowanie bakteryjnej polimerazy RNA zależnej od DNA",
  "streptomycyna – hamowanie syntezy białka"],
 ["pyrazynamid – hamowanie gyrazy w komórce prątka"]),

(46, "UKŁAD POKARMOWY", "Działania niepożądane",
 "Pantoprazol może być przyczyną:",
 ["zmniejszenia wchłaniania witaminy B12 z przewodu pokarmowego",
  "hipergastrynemii", "częstszych złamań kości",
  "zwiększenia liczby zakażeń przewodu pokarmowego"],
 ["hipogastrynemii", "zwiększenia wydzielania kwasu solnego",
  "zwiększenia wchłaniania witaminy B12"]),

(47, "UKŁAD POKARMOWY", "Mechanizm działania",
 "Wskaż właściwe stwierdzenia odnoszące się do metoklopramidu:",
 ["blokuje receptory D-2 w rdzeniu przedłużonym", "zwiększa opróżnianie żołądka"],
 ["pobudza receptory serotoninowe 5-HT3", "działa cholinolitycznie"]),

(48, "UKŁAD POKARMOWY", "Klasyfikacja",
 "Do leków przeczyszczających należy:",
 ["przetwory senesu", "makrogol", "fosforan sodu"],
 ["oktreotyd"]),

(49, "NADCIŚNIENIE TĘTNICZE", "Działania niepożądane",
 "Amlodypina jest przyczyną:",
 ["zaczerwienienia twarzy", "zaparcia", "obrzęków obwodowych"],
 ["bloków serca"]),

(50, "UKŁAD WSPÓŁCZULNY", "Klasyfikacja",
 "Do beta-adrenolityków naczyniorozszerzających należy:",
 ["karwedilol"],
 ["atenolol", "betaksolol", "tymolol"]),

(51, "NIEWYDOLNOŚĆ SERCA", "Działania niepożądane",
 "Podczas stosowania digoksyny może wystąpić:",
 ["blok serca", "ginekomastia", "częstoskurcz węzłowy", "zaburzenie widzenia"],
 ["hipokaliemia", "kaszel", "obrzęk naczynioruchowy"]),

(52, "ZABURZENIA RYTMU SERCA", "Klasyfikacja",
 "Do leków antyarytmicznych należy:",
 ["adenozyna", "amiodaron"],
 ["nimodypina", "alfuzosyna"]),

(53, "LEKI PRZECIWDEPRESYJNE", "Klasyfikacja",
 "Działanie przeciwdepresyjne wykazuje:",
 ["mianseryna", "citalopram", "moklobemid", "imipramina"],
 ["haloperydol", "buspiron", "zolpidem"]),

(54, "LEKI PRZECIWLĘKOWE", "Wskazania",
 "Benzodiazepiny są stosowane w:",
 ["stanach wzmożonego napięcia mięśniowego",
  "leczeniu zespołu abstynencyjnego w chorobie alkoholowej", "premedykacji"],
 ["zatruciu glikozydami naparstnicy"]),

(55, "LEKI PRZECIWPSYCHOTYCZNE", "Działania niepożądane",
 "Złośliwy zespół neuroleptyczny objawia się:",
 ["drżenie włókienkowe", "hipertermia"],
 ["akatyzja", "bradykardia"]),

(56, "ZABURZENIA RYTMU SERCA", "Działania niepożądane",
 "Do leków istotnie wydłużających odstęp QT należy:",
 ["erytromycyna", "astemizol", "moksyfloksacyna", "cyzapryd"],
 ["amoksycylina", "loratadyna", "ranitydyna"]),

(57, "LEKI PRZECIWPSYCHOTYCZNE", "Działania niepożądane",
 "Wskaż właściwe stwierdzenia odnoszące się do klozapiny:",
 ["obniża próg drgawkowy", "powoduje agranulocytozę"],
 ["częściej niż chloropromazyna powoduje zespół parkinsonopodobny",
  "powoduje obrzęk plamki żółtej"]),

(58, "UKŁAD POKARMOWY", "Mechanizm działania",
 "Wskaż właściwe stwierdzenia odnoszące się do ondansetronu:",
 ["jest stosowany w wymiotach przy chemioterapii",
  "jest antagonistą receptorów 5-HT3"],
 ["powoduje hiperprolaktynemię", "podwyższa stężenie trójglicerydów we krwi"]),

(59, "UKŁAD POKARMOWY", "Wskazania",
 "W chorobie lokomocyjnej zastosujesz:",
 ["skopolamina", "dimenhydrynat"],
 ["lorazepam", "aprepitant"]),

(60, "ZNIECZULENIA OGÓLNE", "Sposób podania",
 "Wziewnie w znieczuleniu ogólnym zastosujesz:",
 ["halotan", "podtlenek azotu"],
 ["etomidat", "propofol"]),

(61, "GLIKOKORTYKOSTEROIDY", "Mechanizm działania",
 "Glikokortykosteroidy zmniejszają liczbę:",
 ["monocytów", "limfocytów"],
 ["płytek krwi", "neutrofili"]),

(62, "LEKI PRZECIWHISTAMINOWE", "Wskazania",
 "Leki blokujące receptor H1 pierwszej generacji znalazły zastosowanie:",
 ["w chorobie lokomocyjnej", "w leczeniu zawrotów głowy",
  "w leczeniu dermatologicznych zmian atopowych", "we wstrząsie anafilaktycznym"],
 ["w leczeniu astmy oskrzelowej jako leki pierwszego rzutu",
  "w leczeniu jaskry z zamkniętym kątem przesączania",
  "w leczeniu nadciśnienia tętniczego"]),

(63, "NARKOTYCZNE LEKI PRZECIWBÓLOWE", "Działania niepożądane",
 "Po zastosowaniu kodeiny może wystąpić:",
 ["zaparcie", "senność", "zwężenie źrenic"],
 ["kaszel"]),

(64, "UKŁAD ODDECHOWY", "Mechanizm działania",
 "Wybierz prawidłowe połączenie lek – mechanizm działania:",
 ["bromheksyna – aktywacja enzymów hydrolitycznych",
  "dornaza alfa – hydroliza DNA zawartego w plwocinie",
  "montelukast – antagonista receptorów leukotrienów CysLT1"],
 ["kodeina – antagonista receptorów opioidowych"]),

(65, "NIEOPIOIDOWE LEKI PRZECIWBÓLOWE", "Działania niepożądane",
 "Wskaż prawidłowy zestaw lek – działanie niepożądane:",
 ["paracetamol – hepatotoksyczność", "kwas acetylosalicylowy – nefrotoksyczność",
  "indometacyna – zaburzenia neurologiczne", "metamizol – agranulocytoza"],
 ["celekoksyb – hipertermia złośliwa", "ibuprofen – methemoglobinemia",
  "kwas acetylosalicylowy – zespół Cushinga"]),

(66, "FARMAKOKINETYKA", "Farmakokinetyka",
 "Zjawisko samoindukcji:",
 ["to zdolność niektórych leków do przyspieszania własnego metabolizmu",
  "zachodzi podczas długotrwałego stosowania niektórych leków",
  "jest przyczyną powstawania tolerancji na lek",
  "może zwiększać skuteczność terapeutyczną leku, gdy formą aktywną jest metabolit"],
 ["polega na hamowaniu własnego metabolizmu przez lek",
  "występuje wyłącznie po podaniu dożylnym",
  "dotyczy wyłącznie leków wydalanych w postaci niezmienionej"]),

(67, "CUKRZYCA", "Działania niepożądane",
 "Działaniem niepożądanym insuliny jest:",
 ["lipodystrofia poinsulinowa", "obrzęki", "hipertrofia poinsulinowa", "hipoglikemia"],
 ["kwasica ketonowa", "agranulocytoza", "hiperkaliemia"]),

(68, "ANTYKONCEPCJA", "Mechanizm działania",
 "Działanie środków antykoncepcyjnych polega na:",
 ["niedopuszczeniu do owulacji przez zahamowanie wydzielania gonadotropin",
  "zmianie konsystencji śluzu szyjkowego",
  "zaburzeniu perystaltyki jajowodów",
  "spowodowaniu utraty zdolności do kapacytacji przez plemniki"],
 ["pobudzeniu wydzielania gonadotropin",
  "zwiększeniu grubości endometrium",
  "blokowaniu receptorów androgenowych w jądrze"]),

(69, "NIEOPIOIDOWE LEKI PRZECIWBÓLOWE", "Mechanizm działania",
 "Lekiem selektywnie hamującym cyklooksygenazę-2 jest:",
 ["celekoksyb"],
 ["ibuprofen", "nimesulid", "diklofenak"]),

(70, "NARKOTYCZNE LEKI PRZECIWBÓLOWE", "Klasyfikacja",
 "Do 3 stopnia drabiny analgetycznej należy:",
 ["metadon", "morfina"],
 ["tramadol", "kodeina"]),

(71, "NIEOPIOIDOWE LEKI PRZECIWBÓLOWE", "Mechanizm działania",
 "COX-1 jest nieodwracalnie blokowana przez:",
 ["kwas acetylosalicylowy"],
 ["metamizol", "naproksen", "piroksykam"]),

(72, "LEKI PRZECIWWIRUSOWE", "Wskazania",
 "W leczeniu opryszczki zastosujesz:",
 ["acyklowir"],
 ["gancyklowir", "oseltamiwir", "walgancyklowir"]),

(73, "LEKI PRZECIWPASOŻYTNICZE", "Wskazania",
 "W profilaktyce malarii zastosujesz:",
 ["proguanil", "meflochina"],
 ["pirymetamina", "chinina"]),

(74, "LEKI HIPOLIPEMIZUJĄCE", "Wskazania",
 "Wskaż właściwe stwierdzenia odnoszące się do atorwastatyny:",
 ["jest stosowana w zespole metabolicznym",
  "podawana jest chorym po udarze niedokrwiennym mózgu"],
 ["indukuje reduktazę HMG-CoA", "obniża stężenie frakcji HDL"]),

(75, "BÓLE GŁOWY", "Mechanizm działania",
 "Wskaż właściwe stwierdzenia odnoszące się do ryzatryptanu:",
 ["jest agonistą receptorów 5-HT1B/1D",
  "może być podawany łącznie z niesteroidowym lekiem przeciwzapalnym",
  "jest przeciwwskazany u osób z chorobą wieńcową",
  "jest podawany doustnie"],
 ["jest antagonistą receptorów 5-HT3",
  "jest lekiem pierwszego rzutu w profilaktyce migreny",
  "może być podawany łącznie z inhibitorami MAO"]),

(76, "CHOROBA PARKINSONA", "Klasyfikacja",
 "Agonistą receptora dopaminergicznego jest:",
 ["pergolid", "kabergolina"],
 ["oksybutynina", "selegilina"]),

(77, "PADACZKA", "Klasyfikacja",
 "Działanie przeciwdrgawkowe wykazuje:",
 ["tiagabina", "kwas walproinowy", "gabapentyna", "lamotrygina"],
 ["flumazenil", "teofilina", "bupropion"]),

(78, "STANY NAGŁE", "Działania niepożądane",
 "Wskaż właściwe stwierdzenia odnoszące się do hydroksyetyloskrobi:",
 ["gromadzi się w narządach miąższowych", "powoduje silny świąd"],
 ["jest hepatotoksyczna", "powoduje agranulocytozę"]),

(79, "JASKRA", "Wskazania",
 "W jaskrze zastosujesz:",
 ["acetazolamid", "tymolol", "pilokarpina"],
 ["solifenacyna"]),

(80, "UKŁAD PRZYWSPÓŁCZULNY", "Wskazania",
 "Neostygmina jest stosowana w:",
 ["nużliwości mięśni", "atonii pooperacyjnej jelit",
  "odwróceniu działania związków wywołujących blok typu niedepolaryzacyjnego"],
 ["zatruciu pestycydami"]),

(81, "LEKI PRZECIWBAKTERYJNE", "Działania niepożądane",
 "Wskaż prawidłowe połączenie lek – działanie niepożądane:",
 ["kotrimoksazol – zespół Stevensa-Johnsona",
  "etambutol – pozagałkowe zapalenie nerwu wzrokowego",
  "ryfampicyna – uszkodzenie wątroby"],
 ["klonidyna – wzrost ciśnienia"]),

(82, "STANY NAGŁE", "Antidotum",
 "Wskaż prawidłowe połączenie lek – leczenie przedawkowania:",
 ["midazolam – flumazenil", "izoniazyd – witamina B6",
  "fentanyl – nalokson", "paracetamol – acetylocysteina"],
 ["morfina – flumazenil", "warfaryna – siarczan protaminy",
  "heparyna – witamina K"]),

(83, "NARKOTYCZNE LEKI PRZECIWBÓLOWE", "Działania niepożądane",
 "Wskaż właściwe stwierdzenia odnoszące się do morfiny:",
 ["zwiększa uwalnianie histaminy", "zwęża źrenicę", "działa uspokajająco",
  "wykazuje działanie przeciwkaszlowe"],
 ["rozszerza źrenicę", "pobudza ośrodek oddechowy", "nasila perystaltykę jelit"]),

(84, "LEKI PRZECIWBAKTERYJNE", "Wskazania",
 "W zakażeniu gronkowcem możesz zastosować:",
 ["chinuprystyna/dalfoprystyna", "teikoplanina", "imipenem", "kloksacylina"],
 ["aztreonam", "kolistyna", "kwas nalidyksowy"]),

(85, "LEKI PRZECIWBAKTERYJNE", "Wskazania",
 "W leczeniu zakażeń kości możesz zastosować:",
 ["klindamycyna", "wankomycyna", "chloramfenikol"],
 ["netilmycyna"]),

(86, "BÓLE GŁOWY", "Wskazania",
 "Podaj prawidłowe połączenie lek – wskazanie:",
 ["propranolol – drżenie włókienkowe",
  "karbamazepina – neuralgia nerwu trójdzielnego",
  "sumatryptan – migrena", "oseltamiwir – grypa"],
 ["propranolol – astma oskrzelowa", "karbamazepina – jaskra",
  "sumatryptan – choroba wieńcowa"]),

(87, "LEKI TOKSYCZNE W CIĄŻY", "Ciąża",
 "Lekiem przeciwwskazanym do podawania w ciąży jest:",
 ["symwastatyna", "talidomid", "etretynat", "metotreksat"],
 ["penicylina benzylowa", "paracetamol", "insulina"]),

(88, "CHOROBA NIEDOKRWIENNA SERCA", "Wskazania",
 "W anginie Prinzmetala możesz zastosować:",
 ["nitrogliceryna", "nikardypina"],
 ["metoprolol", "karwedilol"]),

(89, "HEMOSTAZA", "Mechanizm działania",
 "Działanie antyagregacyjne wykazuje:",
 ["klopidogrel", "kwas acetylosalicylowy"],
 ["rywaroksaban", "dabigatran"]),

(90, "HEMOSTAZA", "Przeciwwskazania",
 "Heparyna jest przeciwwskazana u chorych:",
 ["z nadciśnieniem złośliwym", "z małopłytkowością"],
 ["z chorobą nowotworową", "leżących"]),

(91, "HEMOSTAZA", "Działania niepożądane",
 "Działaniem niepożądanym heparyny drobnocząsteczkowej jest:",
 ["krwawienie"],
 ["osteoporoza", "hiperaldosteronizm", "łysienie"]),

(92, "HEMOSTAZA", "Wskazania",
 "Leczenie trombolityczne zastosujesz w:",
 ["zawał mięśnia sercowego", "udar niedokrwienny mózgu"],
 ["obrzęk płuc", "encefalopatia nadciśnieniowa"]),

(93, "NIEOPIOIDOWE LEKI PRZECIWBÓLOWE", "Wskazania",
 "Naproksen wykazuje skuteczność w:",
 ["ból miesiączkowy"],
 ["ból zawałowy", "ból porodowy",
  "profilaktyka udaru niedokrwiennego mózgu", "neuralgia półpaścowa"]),

(94, "HEMOSTAZA", "Wskazania",
 "W krwawieniu o nieznanej przyczynie zastosujesz:",
 ["kwas traneksamowy"],
 ["prasugrel", "iwabradyna", "fondaparynuks", "meloksykam"]),

(95, "FARMAKOKINETYKA", "Farmakokinetyka",
 "Dostępność biologiczna leku przy podaniu doustnym nie zależy od:",
 ["czasu biologicznego półtrwania"],
 ["stopnia jonizacji leku", "wielkości cząsteczki leku", "obecności pokarmu",
  "perystaltyki przewodu pokarmowego"]),

(96, "LEKI PRZECIWBAKTERYJNE", "Wskazania",
 "W leczeniu oparzeń możesz zastosować:",
 ["jodowany poliwidon"],
 ["chlorheksydyna", "nystatyna", "mupirocyna", "fiolet krystaliczny"]),

(97, "LEKI PRZECIWHISTAMINOWE", "Klasyfikacja",
 "Działanie przeciwhistaminowe wykazuje:",
 ["akrywastyna"],
 ["estazolam", "mianseryna", "ksylometazolina", "donepezil"]),

(98, "LEKI MOCZOPĘDNE", "Wskazania",
 "Wskaż właściwe stwierdzenie odnoszące się do mannitolu:",
 ["jest wskazany w obrzęku mózgu"],
 ["jest podawany doustnie", "jest wskazany w obrzęku płuc",
  "podwyższa stężenie potasu", "jest agonistą receptorów alfa"]),

(99, "UKŁAD ODDECHOWY", "Sposób podania",
 "Wziewnie w astmie oskrzelowej zastosujesz:",
 ["beklometazon"],
 ["metyloprednizolon", "triamcynolon", "hydrokortyzon", "ketotifen"]),

(100, "STANY NAGŁE", "Wskazania",
 "Adrenalina jest wskazana w:",
 ["stan astmatyczny", "wstrząs anafilaktyczny",
  "jaskra z otwartym kątem przesączania", "resuscytacja krążeniowo-oddechowa"],
 ["nadciśnienie tętnicze", "przełom tarczycowy", "guz chromochłonny nadnerczy"]),
]

MAX_T, MAX_F = 5, 5
HEADER = (["Section", "Category", "Poziom", "Type", "Question"]
          + [f"True {i}" for i in range(1, MAX_T + 1)]
          + [f"False {i}" for i in range(1, MAX_F + 1)])


def rows():
    for nr, sec, cat, q, trues, falses in Q:
        assert trues, f"pytanie {nr}: brak True"
        assert len(trues) <= MAX_T, f"pytanie {nr}: za duzo True"
        assert len(falses) <= MAX_F, f"pytanie {nr}: za duzo False"
        yield ([sec, cat, "", "testowe", q]
               + trues + [""] * (MAX_T - len(trues))
               + falses + [""] * (MAX_F - len(falses)))


def main():
    data = list(rows())
    assert len(data) == 100, len(data)

    with open("farmakologia_egzamin_2014_2015_OnlyPharms.csv", "w",
              newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(HEADER)
        w.writerows(data)

    wb = Workbook()
    ws = wb.active
    ws.title = "Pytania"
    ws.append(HEADER)
    for r in data:
        ws.append(r)

    head_fill = PatternFill("solid", fgColor="2F5597")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}{len(data) + 1}"

    widths = {"A": 32, "B": 22, "C": 9, "D": 10, "E": 62}
    for col, wdt in widths.items():
        ws.column_dimensions[col].width = wdt
    for i in range(6, len(HEADER) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 34
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save("farmakologia_egzamin_2014_2015_OnlyPharms.xlsx")

    from collections import Counter
    print("wierszy:", len(data))
    print("Trues laczne:", sum(len(t) for _, _, _, _, t, _ in Q))
    print("Falses laczne:", sum(len(f) for _, _, _, _, _, f in Q))
    print("\nSection:")
    for k, v in sorted(Counter(x[1] for x in Q).items()):
        print(f"  {v:3d}  {k}")
    print("\nCategory:")
    for k, v in sorted(Counter(x[2] for x in Q).items()):
        print(f"  {v:3d}  {k}")


if __name__ == "__main__":
    main()
