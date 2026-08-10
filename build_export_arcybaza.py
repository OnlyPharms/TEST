# -*- coding: utf-8 -*-
"""
Buduje plik eksportowy Only Pharms z dokumentu "ARCYBAZA FARMA Z ODP".

Zrodlo: 367-stronicowa skladanka egzaminow i kolokwiow z farmakologii
(kierunek lekarski / IWL), wersja 29.05.2020. Poprawne odpowiedzi sa
w PDF oznaczone KOLOREM PODSWIETLENIA, nie tekstem - odczytane z obiektow
graficznych PDF (zielony 0,255,0 oraz zolty 255,255,0).

Pipeline (skrypty robocze w scratchpadzie):
  extract_arcy.py  - tekst + podswietlenia per znak
  parse2.py        - sekcje / pytania / opcje (dwa style: a.b.c.d oraz 1.2.3.4)
  classify.py      - automatyczne Section/Category ze slownikow JAK_KATEGORYZOWAC

Uklad kolumn: Section, Category, Poziom, Type, Question, True 1..5, False 1..5
plus dwie kolumny informacyjne (Zrodlo, Pewnosc), ktore importer ignoruje.
"""
import csv, json, os, re, sys, unicodedata
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

SCRATCH = ('/tmp/claude-0/-home-user-TEST/'
           '818078ff-c57a-51d2-987e-4f2584083ea6/scratchpad')
sys.path.insert(0, SCRATCH)
import classify as C  # noqa: E402

MAX_T, MAX_F = 5, 5
HEADER = (["Section", "Category", "Poziom", "Type", "Question"]
          + [f"True {i}" for i in range(1, MAX_T + 1)]
          + [f"False {i}" for i in range(1, MAX_F + 1)]
          + ["Zrodlo", "Pewnosc"])


def norm(s):
    return re.sub(r'[^a-z0-9]', '',
                  unicodedata.normalize('NFKD', s.lower())
                  .encode('ascii', 'ignore').decode())


def collect():
    secs = json.load(open(os.path.join(SCRATCH, 'arcy_parsed2.json')))
    seen, out, stats = {}, [], Counter()
    for s in secs:
        for q in s['qs']:
            stats['wszystkie'] += 1
            if not any(o['tag'] in ('G', 'Y') for o in q['opts']):
                stats['bez klucza'] += 1
                continue
            key = norm(q['text']) + '|' + '|'.join(
                sorted(norm(o['text']) for o in q['opts']))
            if key in seen:
                stats['duplikaty'] += 1
                continue
            seen[key] = 1
            q['src'] = s['name']
            out.append(q)
    return out, stats


def build():
    qs, stats = collect()
    rows, dropped = [], []
    noF = 0
    for q in qs:
        trues = [o['text'] for o in q['opts'] if o['tag'] in ('G', 'Y')]
        falses = [o['text'] for o in q['opts'] if o['tag'] not in ('G', 'Y')]
        sec, cat, score = C.classify(q['text'], [o['text'] for o in q['opts']])
        if sec is None:
            dropped.append(q)
            continue
        if not trues:
            dropped.append(q)
            continue
        if not falses:
            noF += 1
        trues, falses = trues[:MAX_T], falses[:MAX_F]
        conf = 'wysoka' if score >= 3 else ('srednia' if score == 2 else 'niska')
        rows.append([sec, cat, 'lekarski', 'testowe', q['text']]
                    + trues + [''] * (MAX_T - len(trues))
                    + falses + [''] * (MAX_F - len(falses))
                    + [q['src'], conf])
    return rows, stats, dropped, noF


def main():
    rows, stats, dropped, noF = build()
    base = 'farmakologia_ARCYBAZA_OnlyPharms'

    with open(base + '.csv', 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(HEADER)
        w.writerows(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Pytania'
    ws.append(HEADER)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='2F5597')
        c.alignment = Alignment(vertical='center')
    ws.freeze_panes = 'F2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(HEADER))}{len(rows) + 1}'
    for col, wdt in {'A': 30, 'B': 21, 'C': 9, 'D': 9, 'E': 60}.items():
        ws.column_dimensions[col].width = wdt
    for i in range(6, len(HEADER) - 1):
        ws.column_dimensions[get_column_letter(i)].width = 32
    ws.column_dimensions[get_column_letter(len(HEADER) - 1)].width = 26
    ws.column_dimensions[get_column_letter(len(HEADER))].width = 10
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical='top', wrap_text=True)
    wb.save(base + '.xlsx')

    print('sparsowane pytania:      ', stats['wszystkie'])
    print('  bez klucza (odrzucone):', stats['bez klucza'])
    print('  duplikaty (odrzucone): ', stats['duplikaty'])
    print('  bez dzialu (odrzucone):', len(dropped))
    print('WYEKSPORTOWANO:          ', len(rows))
    print('  w tym bez dystraktorow (same poprawne):', noF)
    print()
    print('pewnosc klasyfikacji:', dict(Counter(r[-1] for r in rows)))
    print()
    for k, v in Counter(r[0] for r in rows).most_common():
        print(f'  {v:5d}  {k}')


if __name__ == '__main__':
    main()
